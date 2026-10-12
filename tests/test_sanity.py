"""Sanity checks. Run with `python tests/test_sanity.py` (or pytest)."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402

from dwc import config as C, forecast as F  # noqa: E402
from dwc.history_match import from_cube, select_diverse, to_cube  # noqa: E402
from dwc.proxy import to_unit  # noqa: E402


def test_powerlaw_recovers_known_exponent():
    days = np.full(40, 91.25)
    liq = np.full(40, 13000.0)
    x = F._mid_cum(liq, days)
    oil = liq * np.minimum(1, 0.9 * (x / x[0]) ** -1.3)
    _, beta = F.fit_powerlaw_oilcut(oil, liq - oil, days, 9)
    assert abs(beta - 1.3) < 1e-6


def test_forecast_continuous_and_monotonic():
    days = np.full(41, 91.25)
    liq = np.full(41, 13000.0)
    x = F._mid_cum(liq, days)
    oil = liq * np.minimum(1, 0.9 * (x / x[0]) ** -1.3)
    obs = {"oil": oil, "water": liq - oil, "gas": 0.363 * oil, "days": days,
           "cum_anchor": {"Oil": 1e7, "Water": 5e6, "Gas": 3.63e6}}
    df = F.field_forecast(obs, (0.0, 0.05))[0]
    assert len(df) == 80 and df["date"].iloc[0].strftime("%Y-%m-%d") == "2008-04-01"
    assert df["date"].iloc[-1].strftime("%Y-%m-%d") == "2028-01-01"
    for col in ("oil_cum_stb", "water_cum_stb", "gas_cum_mscf"):
        assert np.all(np.diff(df[col].values) > 0)
    # first forecast quarter continues the synthetic power law exactly (anchored on the last point)
    fc_days = F.forecast_days()
    x_next = np.sum(liq * days) + 0.5 * liq[-1] * fc_days[0]
    expected = liq[-1] * 0.9 * (x_next / x[0]) ** -1.3
    assert abs(df["oil_rate_stbd"].iloc[1] / expected - 1) < 1e-6


def test_parameter_transforms_roundtrip():
    rng = np.random.default_rng(0)
    th = from_cube(rng.random((50, 4)))
    assert np.allclose(to_cube(from_cube(to_cube(th))), to_cube(th))
    u = to_unit(th)
    assert u.min() >= -1e-12 and u.max() <= 1 + 1e-12


def test_select_diverse_respects_distance():
    rng = np.random.default_rng(1)
    th = from_cube(rng.random((500, 4)))
    idx = select_diverse(th, rng.random(500), k=10, min_dist=0.2)
    U = to_unit(th[idx])
    d = np.linalg.norm(U[:, None] - U[None], axis=-1) + np.eye(len(idx)) * 9
    assert len(idx) == 10 and d.min() >= 0.2


def test_deliverable_if_present():
    path = C.DELIVERABLE / "09_Template_Deliverable.xlsx"
    if not path.exists():
        print("  (deliverable not built yet - skipped)")
        return
    wb = openpyxl.load_workbook(path, data_only=True)
    assert [f"Case {i}" for i in range(1, 11)] == wb.sheetnames[:10]
    for name in wb.sheetnames[:10]:
        ws = wb[name]
        vals = np.array([[ws.cell(r, c).value for c in (2, 3, 4)] for r in range(2, 82)], float)
        assert not np.isnan(vals).any() and np.all(np.diff(vals, axis=0) > 0)
        gor = (vals[-1, 0] - vals[0, 0]) / (vals[-1, 1] - vals[0, 1])
        assert 0.35 < gor < 0.38, gor


if __name__ == "__main__":
    tests = [v for k, v in dict(globals()).items() if k.startswith("test_")]
    for t in tests:
        t()
        print(f"ok  {t.__name__}")
    print(f"{len(tests)} tests passed")
